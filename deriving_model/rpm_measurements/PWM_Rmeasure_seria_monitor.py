"""
PWM Serial Monitor Reader -> Excel
Διαβάζει RPM από Arduino στο COM8 και τα εξάγει σε xlsx.

Εγκατάσταση εξαρτήσεων:
    pip install pyserial openpyxl
"""

import serial
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import re
from datetime import datetime

# ---------- CONFIG ----------
PORT = 'COM8'
BAUDRATE = 115200
PWM_VALUES = list(range(150, 260, 10))   # 150, 160, ... 250
HIGHLIGHT_PWM = 200                      # κίτρινη στήλη
OUTPUT_FILE = f'rpm_measurements_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
# ----------------------------


def read_serial():
    """Διαβάζει από το serial μέχρι να δει DONE ή να διακοπεί χειροκίνητα."""
    data = {pwm: [] for pwm in PWM_VALUES}
    current_pwm = None

    print(f"[i] Σύνδεση στο {PORT} @ {BAUDRATE} baud ...")
    ser = serial.Serial(PORT, BAUDRATE, timeout=2)
    print(f"[✓] Συνδέθηκα. Περιμένω δεδομένα από το Arduino...")
    print("    (Ctrl+C για πρόωρη διακοπή — θα σωθεί ό,τι έχει μαζευτεί)\n")

    try:
        while True:
            raw = ser.readline()
            if not raw:
                continue
            line = raw.decode('utf-8', errors='ignore').strip()
            if not line:
                continue

            # echo στο terminal για να βλέπεις live
            print(f"  >> {line}")

            # Header τύπου "PWM = 150"
            m = re.match(r'PWM\s*=\s*(\d+)', line, re.IGNORECASE)
            if m:
                pwm = int(m.group(1))
                if pwm in data:
                    current_pwm = pwm
                    print(f"[→] Καταγραφή στη στήλη PWM={pwm}")
                else:
                    current_pwm = None
                continue

            # Διαχωριστικό
            if set(line) <= {'-', '='}:
                continue

            # Τέλος
            if 'DONE' in line.upper():
                print("\n[✓] Όλες οι μετρήσεις ολοκληρώθηκαν.")
                break

            # Προσπάθεια parse RPM (το Arduino στέλνει με τελεία)
            try:
                rpm = float(line.replace(',', '.'))
            except ValueError:
                continue

            if current_pwm is not None:
                data[current_pwm].append(rpm)

    except KeyboardInterrupt:
        print("\n[!] Διακοπή από χρήστη.")
    finally:
        ser.close()

    return data


def save_to_xlsx(data, filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'PWM vs RPM'

    # Στυλ
    bold = Font(bold=True)
    center = Alignment(horizontal='center', vertical='center')
    yellow = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    thin = Side(border_style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Header row: A1 = "PWM", B1..L1 = 150..250
    ws.cell(row=1, column=1, value='PWM')
    for col_idx, pwm in enumerate(PWM_VALUES, start=2):
        ws.cell(row=1, column=col_idx, value=pwm)

    # Πόσες γραμμές δεδομένων (max όλων των στηλών)
    max_samples = max((len(v) for v in data.values()), default=0)

    # Data rows
    for i in range(max_samples):
        row = i + 2
        ws.cell(row=row, column=1, value=f'#{i + 1}')
        for col_idx, pwm in enumerate(PWM_VALUES, start=2):
            values = data[pwm]
            if i < len(values):
                cell = ws.cell(row=row, column=col_idx, value=round(values[i], 1))
                cell.number_format = '0.0'
            else:
                ws.cell(row=row, column=col_idx, value=0)

    # Format όλα τα κελιά
    last_col = len(PWM_VALUES) + 1
    last_row = max_samples + 1
    for r in range(1, last_row + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = center
            cell.border = border
            if r == 1 or c == 1:
                cell.font = bold

    # Κίτρινο highlight στη στήλη του HIGHLIGHT_PWM
    if HIGHLIGHT_PWM in PWM_VALUES:
        hcol = PWM_VALUES.index(HIGHLIGHT_PWM) + 2
        for r in range(1, last_row + 1):
            ws.cell(row=r, column=hcol).fill = yellow

    # Πλάτη στηλών
    ws.column_dimensions['A'].width = 7
    for c in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 8

    # Freeze πρώτη γραμμή & πρώτη στήλη
    ws.freeze_panes = 'B2'

    wb.save(filename)
    print(f"[✓] Αποθηκεύτηκε στο: {filename}")

    # Σύνοψη
    print("\nΣύνοψη μετρήσεων:")
    for pwm in PWM_VALUES:
        print(f"  PWM={pwm}: {len(data[pwm])} τιμές")


def main():
    try:
        data = read_serial()
    except serial.SerialException as e:
        print(f"[✗] Σφάλμα serial: {e}")
        print("    Έλεγξε ότι το Arduino είναι συνδεδεμένο στο σωστό COM port")
        print("    και ότι δεν είναι ανοιχτός ταυτόχρονα ο Serial Monitor του IDE.")
        return

    if not any(data.values()):
        print("[!] Δεν συλλέχθηκε καμία τιμή — το αρχείο δεν θα δημιουργηθεί.")
        return

    save_to_xlsx(data, OUTPUT_FILE)


if __name__ == '__main__':
    main()